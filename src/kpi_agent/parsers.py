from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class ParsedDocument:
    source_path: str
    text: str


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    # pypdf is a light dependency and works offline.
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    # python-docx is a light dependency and works offline.
    import docx  # type: ignore

    d = docx.Document(str(path))
    parts: list[str] = []
    for p in d.paragraphs:
        if p.text:
            parts.append(p.text)
    # Tables often contain requirements; include them as well.
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                t = (cell.text or "").strip()
                if t:
                    parts.append(t)
    return "\n".join(parts)


def parse_document(path: str | Path) -> ParsedDocument:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Input file not found: {p}")

    suffix = p.suffix.lower()
    if suffix in {".txt", ".md"}:
        text = _read_text_file(p)
    elif suffix == ".pdf":
        text = _read_pdf(p)
    elif suffix == ".docx":
        text = _read_docx(p)
    else:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Supported: .pdf, .docx, .txt, .md"
        )

    text = (text or "").strip()
    return ParsedDocument(source_path=str(p), text=text)

